"""Đọc PDF / DOCX / TXT thành các phần văn bản, mỗi phần kèm vị trí (trang, mục) để trích dẫn."""

import unicodedata
from dataclasses import dataclass
from pathlib import Path

import docx
import pymupdf
from docx.table import Table

SUPPORTED_TYPES = ("pdf", "docx", "txt", "md")


@dataclass
class Section:
    text: str
    location: str  # "trang 5", "mục 2. Chính sách bảo hành"... hoặc "" nếu không xác định được


@dataclass
class Document:
    sections: list[Section]
    pages: int | None = None  # chỉ PDF mới có số trang
    title: str = ""  # tiêu đề trong front matter (trang web do crawler tải về), rỗng nếu không có


def load_document(path: Path) -> Document:
    ext = path.suffix.lower().lstrip(".")
    if ext == "pdf":
        return _load_pdf(path)
    if ext == "docx":
        return Document(_load_docx(path))
    if ext in ("txt", "md"):
        return _load_text(path)
    raise ValueError(f"Chưa hỗ trợ định dạng .{ext}")


def _load_pdf(path: Path) -> Document:
    """Mỗi trang là một phần, nên chunk không bao giờ vắt qua 2 trang và trích dẫn đúng số trang."""
    with pymupdf.open(path) as pdf:
        page_count = pdf.page_count
        sections = [Section(_clean(page.get_text()), f"trang {i}") for i, page in enumerate(pdf, start=1)]
    sections = [s for s in sections if s.text]
    if page_count and not sections:
        raise ValueError("PDF không có lớp chữ (có thể là ảnh scan). Cần chạy OCR trước khi nạp.")
    return Document(sections, pages=page_count)


def _load_docx(path: Path) -> list[Section]:
    """DOCX không lưu số trang, nên chia theo tiêu đề (Heading) để trích dẫn theo mục."""
    sections: list[Section] = []
    heading, lines = "", []

    def flush():
        text = _clean("\n".join(lines))
        if text:
            sections.append(Section(text, f"mục {heading}" if heading else ""))
        lines.clear()

    for block in docx.Document(path).iter_inner_content():  # đoạn văn và bảng theo đúng thứ tự
        if isinstance(block, Table):
            lines.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in block.rows)
            continue
        text = block.text.strip()
        if not text:
            continue
        style = block.style.name.lower() if block.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            heading = text
        lines.append(text)
    flush()
    return sections


def _load_text(path: Path) -> Document:
    """TXT/Markdown. File do crawler tạo có front matter (title, url) ở đầu: URL dùng làm nguồn trích dẫn."""
    text = _clean(path.read_text(encoding="utf-8-sig", errors="replace"))
    fields = {}
    if text.startswith("---\n") and "\n---\n" in text:
        header, _, text = text[4:].partition("\n---\n")
        fields = dict(line.split(": ", 1) for line in header.splitlines() if ": " in line)
    text = text.strip()
    sections = [Section(text, fields.get("url", ""))] if text else []
    return Document(sections, title=fields.get("title", ""))


def _clean(text: str) -> str:
    """Chuẩn hoá để tìm kiếm ổn định: chữ Việt về dạng dựng sẵn (NFC) — cùng chữ "ệ" có thể được lưu
    thành 1 hoặc 3 ký tự tuỳ phần mềm soạn thảo; dấu cách không ngắt (\\xa0) hay gặp trong PDF thành dấu cách
    thường; bỏ gạch nối mềm (\\xad)."""
    return unicodedata.normalize("NFC", text).replace("\xa0", " ").replace("\xad", "").strip()
